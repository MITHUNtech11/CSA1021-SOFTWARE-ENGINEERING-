import sys
import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_deck(output_path):
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Classic Academic / Clean Minimal
    NAVY = RGBColor(26, 54, 93)         # #1A365D - Deep Academic Navy
    BLUE = RGBColor(37, 99, 235)        # #2563EB - Accent Royal Blue
    SLATE_DARK = RGBColor(15, 23, 42)   # #0F172A - Headings / Charcoal
    SLATE_BODY = RGBColor(51, 65, 85)   # #334155 - Body Text
    SLATE_MUTED = RGBColor(100, 116, 139) # #64748B - Secondary / Metadata
    BORDER_COLOR = RGBColor(226, 232, 240) # #E2E8F0 - Clean borders
    BG_CARD = RGBColor(248, 250, 252)   # #F8FAFC - Card fill
    BG_BADGE = RGBColor(238, 242, 255)  # #EEF2FF - Badge pill fill
    BADGE_TEXT = RGBColor(49, 46, 129)  # #312E81 - Badge text
    WHITE = RGBColor(255, 255, 255)
    RED_ACCENT = RGBColor(225, 29, 72)  # #E11D48 - Highlight / Warning
    GREEN_ACCENT = RGBColor(16, 185, 129) # #10B981 - Success

    FONT_FAMILY = "Calibri"

    def add_header(slide, badge_text, title_text, subtitle_text):
        # Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(4.5), Inches(0.32))
        badge.fill.solid()
        badge.fill.fore_color.rgb = BG_BADGE
        badge.line.color.rgb = RGBColor(199, 210, 254)
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = Inches(0.12)
        tf_b.margin_top = Inches(0.04)
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.name = FONT_FAMILY
        p_b.font.size = Pt(9.5)
        p_b.font.bold = True
        p_b.font.color.rgb = BADGE_TEXT

        # Title & Subtitle Box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_t = tf.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY

        p_s = tf.add_paragraph()
        p_s.text = subtitle_text
        p_s.font.name = FONT_FAMILY
        p_s.font.size = Pt(12)
        p_s.font.color.rgb = SLATE_MUTED
        p_s.space_before = Pt(3)

        # Header Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), Inches(11.73), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_COLOR
        line.line.color.rgb = BORDER_COLOR

    def add_footer(slide, current_slide, total_slides=8):
        # Footer text left
        tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(8.0), Inches(0.3))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True
        tf_l.margin_left = tf_l.margin_top = 0
        p_l = tf_l.paragraphs[0]
        p_l.text = "CSA1021 - Software Engineering | Capstone Project 2026 | Saveetha School of Engineering (SIMATS)"
        p_l.font.name = FONT_FAMILY
        p_l.font.size = Pt(9)
        p_l.font.color.rgb = SLATE_MUTED

        # Footer slide number right
        tb_r = slide.shapes.add_textbox(Inches(10.5), Inches(7.0), Inches(2.0), Inches(0.3))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True
        tf_r.margin_right = tf_r.margin_top = 0
        p_r = tf_r.paragraphs[0]
        p_r.text = f"Slide {current_slide} of {total_slides}"
        p_r.alignment = PP_ALIGN.RIGHT
        p_r.font.name = FONT_FAMILY
        p_r.font.size = Pt(9)
        p_r.font.color.rgb = SLATE_MUTED

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)

    # Decorative Top Accent Bar
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = NAVY
    top_bar.line.fill.background()

    # Department Badge
    dept_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.7), Inches(11.333), Inches(0.42))
    dept_badge.fill.solid()
    dept_badge.fill.fore_color.rgb = BG_BADGE
    dept_badge.line.color.rgb = RGBColor(199, 210, 254)
    dept_badge.line.width = Pt(1)
    tf_db = dept_badge.text_frame
    tf_db.margin_top = Inches(0.08)
    p_db = tf_db.paragraphs[0]
    p_db.text = "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING | SAVEETHA SCHOOL OF ENGINEERING (SIMATS)"
    p_db.alignment = PP_ALIGN.CENTER
    p_db.font.name = FONT_FAMILY
    p_db.font.size = Pt(11)
    p_db.font.bold = True
    p_db.font.color.rgb = BADGE_TEXT

    # Main Title Box
    title_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(11.333), Inches(2.3))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t1 = tf_t.paragraphs[0]
    p_t1.text = "EventFlow: Preventing Double-Booking in Campus Venue Scheduling Using the Interval Scheduling Algorithm"
    p_t1.alignment = PP_ALIGN.CENTER
    p_t1.font.name = FONT_FAMILY
    p_t1.font.size = Pt(28)
    p_t1.font.bold = True
    p_t1.font.color.rgb = NAVY
    p_t1.space_after = Pt(12)

    p_t2 = tf_t.add_paragraph()
    p_t2.text = "A Centralized, Real-Time Campus Event & Session Coordination System"
    p_t2.alignment = PP_ALIGN.CENTER
    p_t2.font.name = FONT_FAMILY
    p_t2.font.size = Pt(15)
    p_t2.font.bold = True
    p_t2.font.color.rgb = BLUE

    # Mid Accent Divider
    mid_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(3.8), Inches(4.333), Inches(0.03))
    mid_line.fill.solid()
    mid_line.fill.fore_color.rgb = BLUE
    mid_line.line.fill.background()

    # Two Metadata Cards
    # Left Card: Project Metadata
    c_left = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.15), Inches(5.2), Inches(2.4))
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = BG_CARD
    c_left.line.color.rgb = BORDER_COLOR
    c_left.line.width = Pt(1)
    tf_cl = c_left.text_frame
    tf_cl.margin_left = tf_cl.margin_top = tf_cl.margin_right = Inches(0.25)
    tf_cl.word_wrap = True

    p = tf_cl.paragraphs[0]
    p.text = "ACADEMIC & COURSE METADATA"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(8)

    meta_items = [
        ("Course Code & Title:", "CSA1021 - Software Engineering"),
        ("Curriculum Scope:", "Capstone Project Assessment (2026)"),
        ("Project Platform:", "EventFlow (FastAPI + Supabase / SQLite + React)"),
        ("Core Focus Area:", "Interval Scheduling & Temporal Concurrency"),
    ]
    for lbl, val in meta_items:
        p_m = tf_cl.add_paragraph()
        p_m.font.name = FONT_FAMILY
        p_m.font.size = Pt(10)
        p_m.space_after = Pt(4)
        run_lbl = p_m.add_run()
        run_lbl.text = f"{lbl} "
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = SLATE_DARK
        run_val = p_m.add_run()
        run_val.text = val
        run_val.font.color.rgb = SLATE_BODY

    # Right Card: Presenters
    c_right = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(4.15), Inches(5.2), Inches(2.4))
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = BG_CARD
    c_right.line.color.rgb = BORDER_COLOR
    c_right.line.width = Pt(1)
    tf_cr = c_right.text_frame
    tf_cr.margin_left = tf_cr.margin_top = tf_cr.margin_right = Inches(0.25)
    tf_cr.word_wrap = True

    p = tf_cr.paragraphs[0]
    p.text = "PROJECT TEAM & AFFILIATION"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(8)

    team_items = [
        ("Student Lead:", "S. Mithun Senthil (Reg. No: 192324056)"),
        ("Team Member:", "[Teammate Name] (Reg. No: [Teammate Reg No])"),
        ("Department:", "Computer Science and Engineering"),
        ("Institution:", "Saveetha School of Engineering, SIMATS"),
    ]
    for lbl, val in team_items:
        p_m = tf_cr.add_paragraph()
        p_m.font.name = FONT_FAMILY
        p_m.font.size = Pt(10)
        p_m.space_after = Pt(4)
        run_lbl = p_m.add_run()
        run_lbl.text = f"{lbl} "
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = SLATE_DARK
        run_val = p_m.add_run()
        run_val.text = val
        run_val.font.color.rgb = SLATE_BODY

    # Bottom Footer
    add_footer(s1, 1)

    # ==========================================
    # SLIDE 2: INTRODUCTION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Context & Campus Operations", "Introduction: Campus Venue Scheduling Dynamics", 
               "Operational realities and algorithmic challenges of multi-department university event coordination")
    add_footer(s2, 2)

    # 3 Structured Horizontal Cards
    card_data_s2 = [
        ("1. Dynamic University Ecosystem & Multi-Event Scale",
         "Modern academic campuses host dozens of concurrent co-curricular and extracurricular programs each semester—including 36-hour hackathons, technical symposiums, research workshops, guest lectures, cultural galas, and sports tournaments.\n\n"
         "These events coordinate thousands of students, faculty mentors, and specialized campus venues across distinct academic departments operating autonomously.",
         BG_CARD),
        ("2. The Shared Resource & Double-Booking Bottleneck",
         "Colleges operate with finite high-demand facilities (Main Auditorium, Tech Seminar Halls, Smart Labs). Organizers frequently request the exact same hall for identical dates and overlapping times—often within seconds of each other during peak festival scheduling.\n\n"
         "Without a centralized, mathematically verified reservation pipeline, multiple bookings are confirmed for a single venue: the classic double-booking failure.",
         BG_CARD),
        ("3. Computational Interval Model & EventFlow Integration",
         "In computer science, each venue booking is formulated as a discrete continuous time interval [start, end) allocated to a specific physical resource ID.\n\n"
         "EventFlow solves this structural failure by embedding rigorous interval scheduling directly into a centralized campus event coordination platform, providing conflict-free scheduling alongside digital QR passes and session navigation.",
         RGBColor(241, 245, 249))
    ]

    left_positions = [Inches(0.8), Inches(4.8), Inches(8.8)]
    for i, (head, body, bg_col) in enumerate(card_data_s2):
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_positions[i], Inches(2.05), Inches(3.73), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_col
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = head
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        p0.space_after = Pt(12)

        p1 = tf.add_paragraph()
        p1.text = body
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(11)
        p1.font.color.rgb = SLATE_BODY
        p1.line_spacing = 1.25

    # ==========================================
    # SLIDE 3: OBJECTIVE
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Algorithmic & System Goals", "Project Objectives", 
               "Core technical benchmarks and architectural outcomes established for EventFlow")
    add_footer(s3, 3)

    obj_cards = [
        ("Objective 1: Strict Temporal Collision Prevention",
         "Detect and preemptively block overlapping bookings for the same campus venue before database confirmation. Treat bookings as formal time intervals [start, end) and evaluate intersection logic to eliminate complete, partial, and multi-round clashes.",
         BLUE),
        ("Objective 2: Optimal Venue Throughput via Interval Scheduling",
         "Apply the classic greedy Earliest Finish Time (EFT) interval scheduling algorithm for batch or competing reservation windows, accepting the largest possible number of mutually compatible events per campus venue in polynomial time O(M log M).",
         NAVY),
        ("Objective 3: Real-Time High-Concurrency Safety",
         "Enforce atomic transactional locking and sub-second binary search verification (O(log N)) to guarantee that simultaneous multi-user requests cannot bypass validation or induce Time-of-Check to Time-of-Use (TOCTOU) race condition double-bookings.",
         SLATE_DARK),
        ("Objective 4: Intelligent Rejection Feedback & Alternatives",
         "Provide instantaneous, transparent rejection feedback with root-cause conflict metrics. Automatically compute and suggest nearest vacant idle time slots and recommend alternative campus venues possessing equivalent room capacity.",
         BLUE),
    ]

    coords_s3 = [
        (Inches(0.8), Inches(2.05)),
        (Inches(6.8), Inches(2.05)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5))
    ]

    for i, (head, body, acc_col) in enumerate(obj_cards):
        x, y = coords_s3[i]
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.73), Inches(2.25))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        # Left color strip
        strip = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), Inches(2.25))
        strip.fill.solid()
        strip.fill.fore_color.rgb = acc_col
        strip.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.2)

        p0 = tf.paragraphs[0]
        p0.text = head
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        p0.space_after = Pt(8)

        p1 = tf.add_paragraph()
        p1.text = body
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = SLATE_BODY
        p1.line_spacing = 1.2

    # ==========================================
    # SLIDE 4: ABSTRACT
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Executive Summary", "Abstract", 
               "Algorithmic framework and concurrency-safe implementation of EventFlow")
    add_footer(s4, 4)

    # Top Narrative Box
    top_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.05), Inches(11.73), Inches(2.1))
    top_box.fill.solid()
    top_box.fill.fore_color.rgb = BG_CARD
    top_box.line.color.rgb = BORDER_COLOR
    top_box.line.width = Pt(1)
    tf_top = top_box.text_frame
    tf_top.word_wrap = True
    tf_top.margin_left = tf_top.margin_right = tf_top.margin_top = Inches(0.25)

    p_abs = tf_top.paragraphs[0]
    p_abs.text = "In centralized campus event management, double-booking represents a critical architectural failure that disrupts schedules, damages administrative credibility, and strands hundreds of participants. Traditional systems rely on manual oversight or coarse date-level validations that fail to capture partial temporal intersections. This project presents a mathematically rigorous, interval scheduling framework integrated into EventFlow—a modern campus event coordination platform.\n\n" \
                 "By treating every reservation as a half-open temporal interval [s_i, e_i) bound to a discrete venue resource V_k, the system executes real-time collision detection via binary search neighbor evaluation (O(log N)) and employs greedy Earliest Finish Time (EFT) optimization for competing reservation batches. Built with a FastAPI backend and dual-mode persistence (SQLite WAL / Supabase PostgreSQL), all validations and commits execute inside atomic transactions, eliminating Time-of-Check to Time-of-Use (TOCTOU) race conditions."
    p_abs.font.name = FONT_FAMILY
    p_abs.font.size = Pt(11)
    p_abs.font.color.rgb = SLATE_BODY
    p_abs.line_spacing = 1.25

    # 3 Summary Cards below
    s4_pillars = [
        ("Mathematical Formulation",
         "Formal half-open intervals [start, end) model time slots without false boundary conflicts, permitting back-to-back events while strictly blocking overlapping subsets.",
         BLUE),
        ("Concurrency & ACID Integrity",
         "Transactional database locks (WAL mode / PostgreSQL ACID) guarantee atomicity, preventing simultaneous multi-user submissions from bypassing availability checks.",
         NAVY),
        ("Platform Verification & Results",
         "Achieves guaranteed zero double-bookings, maximal facility throughput, and automated alternative slot suggestions, fully verified across 28 automated Pytest suites.",
         GREEN_ACCENT)
    ]

    p_lefts = [Inches(0.8), Inches(4.8), Inches(8.8)]
    for i, (head, body, col) in enumerate(s4_pillars):
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, p_lefts[i], Inches(4.35), Inches(3.73), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(241, 245, 249)
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        # Top accent bar on card
        bar = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, p_lefts[i], Inches(4.35), Inches(3.73), Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)

        p0 = tf.paragraphs[0]
        p0.text = head
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        p0.space_after = Pt(6)

        p1 = tf.add_paragraph()
        p1.text = body
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(10)
        p1.font.color.rgb = SLATE_BODY
        p1.line_spacing = 1.2

    # ==========================================
    # SLIDE 5: PROBLEM STATEMENT
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "System Vulnerabilities", "Problem Statement: Vulnerabilities in Traditional Scheduling", 
               "Why conventional institutional booking workflows fail under concurrent campus event demands")
    add_footer(s5, 5)

    prob_cards = [
        ("1. Coarse Granularity & Partial Overlap Blindspots",
         "Simple date-only checks (WHERE date = ?) fail to evaluate temporal ranges, missing partial overlaps (e.g., Event A from 10:00–12:30 colliding with Event B from 11:30–14:00) or complete interval envelopment (09:00–17:00 engulfing 13:00–15:00).",
         RED_ACCENT),
        ("2. Concurrency Vulnerabilities & TOCTOU Race Conditions",
         "When two coordinators request the same venue at 10:00:00 AM, unsynchronized backends execute the read check in parallel. Both threads find the room vacant, both pass validation, and both write INSERT records—guaranteeing an unmitigated double-booking.",
         RED_ACCENT),
        ("3. Multi-Round / Multi-Session Cascade Clashes",
         "Multi-day symposiums and hackathons feature distinct sequential rounds (Screening Quiz, Mentoring, Semifinals, Final Pitching) distributed across diverse campus rooms. A single unhandled collision in Round 2 paralyzes the entire event pipeline.",
         RED_ACCENT),
        ("4. Severe Operational & Reputational Repercussions",
         "Double-booked venues force last-minute emergency relocations, delayed hackathon tracks, microphone/AV resource conflicts, and significant loss of administrative trust between student organizers, department heads, and attendees.",
         RED_ACCENT)
    ]

    coords_s5 = [
        (Inches(0.8), Inches(2.05)),
        (Inches(6.8), Inches(2.05)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5))
    ]

    for i, (head, body, acc_col) in enumerate(prob_cards):
        x, y = coords_s5[i]
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.73), Inches(2.25))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        # Left color strip
        strip = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), Inches(2.25))
        strip.fill.solid()
        strip.fill.fore_color.rgb = acc_col
        strip.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.2)

        p0 = tf.paragraphs[0]
        p0.text = head
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = SLATE_DARK
        p0.space_after = Pt(8)

        p1 = tf.add_paragraph()
        p1.text = body
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = SLATE_BODY
        p1.line_spacing = 1.2

    # ==========================================
    # SLIDE 6: LITERATURE SURVEY (2025-2026)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "State of the Art (2025-2026)", "Literature Survey: Contemporary Scheduling Systems", 
               "Comparative analysis of recent peer-reviewed approaches in timetable and venue scheduling")
    add_footer(s6, 6)

    # Academic Table: 4 rows + header, 4 columns
    rows = 5
    cols = 4
    t_left = Inches(0.8)
    t_top = Inches(2.05)
    t_width = Inches(11.73)
    t_height = Inches(4.7)

    table_shape = s6.shapes.add_table(rows, cols, t_left, t_top, t_width, t_height)
    tbl = table_shape.table

    # Column widths
    tbl.columns[0].width = Inches(3.2)
    tbl.columns[1].width = Inches(2.5)
    tbl.columns[2].width = Inches(2.8)
    tbl.columns[3].width = Inches(3.23)

    headers = [
        "Paper Title & Source (Year)",
        "Core Methodology",
        "Key Strengths",
        "Inherent Limitations & EventFlow Gap"
    ]
    for c_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        tf = cell.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.08)
        p = tf.paragraphs[0]
        p.text = h_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE

    lit_data = [
        (
            "A Course Scheduling System Based on an Improved Genetic Algorithm\n(NHSJS Reports, 2025)",
            "Improved Genetic Algorithm (GA) with adaptive mutation rates and penalty weighting for teacher/room clashes.",
            "Effectively optimizes complex multi-constraint static academic timetables across departments.",
            "Tested only on small static batches; heavy O(G · P) computational overhead; non-viable for instantaneous web booking requests."
        ),
        (
            "Automated Class Scheduling: Ensuring Conflict-Free & Optimized Timetables\n(Khulna Univ. Studies, Dec 2025)",
            "Heuristic-driven automated timetable generator optimizing room and teacher allocations.",
            "Produces clash-free class routines in pre-planned, fixed semester configurations.",
            "Assumes static offline constraints; lacks dynamic ad-hoc request handling and concurrent web transaction safety."
        ),
        (
            "Explicit AI-Based Timetable Generator using Genetic Algorithm\n(Open J. of Applied Sciences, 2025)",
            "AI-assisted Genetic Algorithm optimizing multidimensional lecturer, classroom, and student vectors.",
            "High flexibility in handling soft/hard educational constraints and room preferences.",
            "Stochastic and non-deterministic; requires hundreds of generations; latency prevents sub-second user API responses."
        ),
        (
            "Isabela State University Class Scheduling Algorithm (ISUCSA)\n(J. Inf. Systems Eng. & Mgmt, 2025)",
            "Hybrid evolutionary model combining Genetic Algorithms with a deterministic greedy assignment phase.",
            "Accelerates convergence compared to pure GA; effectively assigns room capacities.",
            "Designed strictly for annual curriculum planning; zero provision for real-time web concurrency or atomic interval locking."
        )
    ]

    for r_idx, row_values in enumerate(lit_data, start=1):
        bg_col = BG_CARD if r_idx % 2 == 1 else WHITE
        for c_idx, val in enumerate(row_values):
            cell = tbl.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_col
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_top = Inches(0.08)
            tf.margin_left = Inches(0.12)
            tf.margin_right = Inches(0.12)
            p = tf.paragraphs[0]
            p.text = val
            p.font.name = FONT_FAMILY
            p.font.size = Pt(9.5)
            p.font.color.rgb = SLATE_DARK if c_idx == 0 else SLATE_BODY
            if c_idx == 0:
                p.font.bold = True

    # ==========================================
    # SLIDE 7: EXISTING METHODS AND DRAWBACKS
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Comparative Deficiency Analysis", "Existing Methods vs. Algorithmic Bottlenecks", 
               "Systematic critique of conventional booking mechanisms and scheduling heuristics")
    add_footer(s7, 7)

    drawback_cards = [
        ("1. Manual Spreadsheets & Paper Registers",
         "Institutional Workflow: Coordinators log requests via Google Forms or Excel; staff manually check calendars.\n"
         "Critical Flaw: Latent conflict discovery—clashes are identified only after both parties receive confirmation. Sluggish and highly error-prone under high event volume.",
         RED_ACCENT),
        ("2. Coarse Date-Level Database Filters",
         "Institutional Workflow: Basic SQL validation checking WHERE date = ? AND venue = ?.\n"
         "Critical Flaw: Coarse temporal resolution—blind to hour/minute intervals. Either rejects non-overlapping events occurring on the same date or admits destructive intra-day overlaps.",
         RED_ACCENT),
        ("3. Genetic Algorithms & Heuristics (Literature)",
         "Institutional Workflow: Iterative multi-generational population optimization (reviewed in Slide 6).\n"
         "Critical Flaw: Massive computational overhead O(G · P · C) with stochastic latency (seconds to minutes). Unsuitable for interactive, real-time web booking APIs.",
         SLATE_DARK),
        ("4. Unsynchronized 'Check-Then-Save' Architectures",
         "Institutional Workflow: Standard web endpoints executing SELECT availability followed by separate INSERT.\n"
         "Critical Flaw: Vulnerable to Time-of-Check to Time-of-Use (TOCTOU) race conditions—simultaneous client requests both evaluate as vacant before either commits.",
         RED_ACCENT),
        ("5. Naive Pairwise Interval Comparisons",
         "Institutional Workflow: Comparing every incoming reservation linearly against all historical records.\n"
         "Critical Flaw: Quadratic O(N^2) scalability degradation as the semester booking database grows, causing severe latency spikes during peak festival registration windows.",
         SLATE_DARK)
    ]

    card_y = Inches(2.05)
    card_h = Inches(0.88)
    card_gap = Inches(0.08)

    for i, (title_t, desc_t, stripe_col) in enumerate(drawback_cards):
        y_pos = card_y + i * (card_h + card_gap)
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.73), card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        # Left color stripe
        stripe = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y_pos, Inches(0.1), card_h)
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = stripe_col
        stripe.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.08)
        tf.margin_right = Inches(0.2)

        p0 = tf.paragraphs[0]
        p0.text = title_t
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(11.5)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        p0.space_after = Pt(2)

        p1 = tf.add_paragraph()
        p1.text = desc_t
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(9.5)
        p1.font.color.rgb = SLATE_BODY
        p1.line_spacing = 1.15

    # ==========================================
    # SLIDE 8: PROPOSED METHODOLOGY
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Architectural Framework", "Proposed Methodology: Algorithmic Interval Scheduling & Concurrency Control", 
               "Three-Pillar Architectural Framework Implemented in the EventFlow Platform")
    add_footer(s8, 8)

    # Three-Pillar Vertical Layout
    pillars_s8 = [
        (
            "Pillar 1: Formal Interval Modeling & Half-Open Algebra",
            [
                ("Resource Mapping:", "Each venue V_k is formulated as an independent, continuous time resource timeline."),
                ("Half-Open Interval:", "Every booking request is represented as I_i = [s_i, e_i), where s_i is start time and e_i is end time."),
                ("Back-to-Back Feasibility:", "Half-open boundaries strictly allow contiguous scheduling ([10:00, 12:00) and [12:00, 14:00) do not collide since 12:00 < 12:00 is False)."),
                ("Overlap Condition:", "Collision(I_new, I_ext) <=> (s_new < e_ext) AND (e_new > s_ext). Any non-zero intersection triggers immediate rejection.")
            ],
            BLUE
        ),
        (
            "Pillar 2: Algorithmic Collision Detection & Optimization",
            [
                ("Single Dynamic Request:", "Active confirmed bookings are kept sorted by start time s_i. Binary search (bisect) locates the insertion index in O(log N) time."),
                ("Neighbor-Only Check:", "Collision is evaluated strictly against adjacent predecessor and successor intervals, eliminating O(N) linear scans."),
                ("Batch Optimization (Greedy EFT):", "For competing reservation batches, intervals are sorted by Earliest Finish Time (e_1 <= e_2 <= ... <= e_M)."),
                ("Provable Optimality:", "Greedily accepting non-overlapping requests guarantees maximum event throughput in O(M log M) time.")
            ],
            NAVY
        ),
        (
            "Pillar 3: Concurrency Engine & Feedback in EventFlow",
            [
                ("ACID Transaction Locking:", "Integrated into FastAPI backend with SQLite WAL mode / Supabase PostgreSQL. Validation and write execute inside an atomic transaction."),
                ("TOCTOU Elimination:", "Venue-level mutual exclusion lock serializes concurrent submissions, preventing race conditions from slipping through."),
                ("Intelligent Feedback:", "When rejected, the engine analyzes schedule gaps to compute the nearest vacant slot in the requested venue."),
                ("Alternative Venue Routing:", "Queries available campus venues with matching capacity, providing immediate alternatives to organizers.")
            ],
            SLATE_DARK
        )
    ]

    col_positions = [Inches(0.8), Inches(4.8), Inches(8.8)]
    for i, (col_title, items, acc_col) in enumerate(pillars_s8):
        c_x = col_positions[i]
        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, Inches(2.05), Inches(3.73), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        # Top Accent Header Bar
        top_bar = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x, Inches(2.05), Inches(3.73), Inches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = acc_col
        top_bar.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.18)

        p0 = tf.paragraphs[0]
        p0.text = col_title
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        p0.space_after = Pt(10)

        for label, desc in items:
            p_item = tf.add_paragraph()
            p_item.font.name = FONT_FAMILY
            p_item.font.size = Pt(9.5)
            p_item.space_after = Pt(8)
            p_item.line_spacing = 1.15

            run_lbl = p_item.add_run()
            run_lbl.text = f"{label} "
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = SLATE_DARK

            run_desc = p_item.add_run()
            run_desc.text = desc
            run_desc.font.color.rgb = SLATE_BODY

    # Save presentation
    prs.save(output_path)
    print(f"Presentation successfully saved to {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "output.pptx"
    create_deck(out)
