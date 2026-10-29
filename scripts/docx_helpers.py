import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=35, bottom=35, left=60, right=60):
    """Set compact inner padding for table cells in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_black_borders(cell):
    """Apply clean, solid black borders to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                          f'<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                          f'<w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                          f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                          f'<w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

def add_heading_1(doc, text):
    """Major Section Headings: 16 pt, Bold, Black."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = BLACK
    return p

def add_heading_2(doc, text):
    """Subheadings: 14 pt, Bold, Black."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = BLACK
    return p

def add_heading_3(doc, text):
    """Minor Subheadings / Labels: 12 pt, Bold, Black."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = BLACK
    return p

def add_body_p(doc, text, bold_prefix=None, italic=False):
    """Normal Body Text: 12 pt, Regular, Black."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = BLACK
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.italic = italic
    run.font.color.rgb = BLACK
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    """Bullet Point: 12 pt, Black with bullet symbol."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.20)
    
    r_bullet = p.add_run("•  ")
    r_bullet.font.name = 'Calibri'
    r_bullet.font.size = Pt(12)
    r_bullet.font.bold = True
    r_bullet.font.color.rgb = BLACK
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = BLACK
        
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.color.rgb = BLACK
    return p

def format_black_lined_table(tbl, col_widths, headers, data_rows, data_font_size=Pt(9.5)):
    """
    Format a clean, unshaded table with solid black outline borders and black text.
    No color fills or background shading.
    """
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = title
        set_cell_margins(cell, top=45, bottom=45, left=60, right=60)
        set_black_borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = BLACK

    # Data Rows
    for r_idx, row_data in enumerate(data_rows):
        row = tbl.add_row()
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_margins(cell, top=35, bottom=35, left=55, right=55)
            set_black_borders(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.08
            run = p.runs[0] if p.runs else p.add_run(val)
            run.font.name = 'Calibri'
            run.font.size = data_font_size
            run.font.color.rgb = BLACK
            if c_idx == 0:
                run.font.bold = True

    # Column widths
    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)
